"""Offline checks for running-run export, workbook display, and score aggregation."""
import contextlib
import csv
import importlib.util
import io
import json
import os
from pathlib import Path
import sys
import tempfile
from types import SimpleNamespace
import unittest
from unittest.mock import Mock, patch

from openpyxl import load_workbook

import schedule_experiments as scheduler


def load_script(name):
    spec = importlib.util.spec_from_file_location(name, Path(__file__).with_name(name + '.py'))
    module = importlib.util.module_from_spec(spec)
    with patch.object(sys, 'path', list(sys.path)):
        spec.loader.exec_module(module)
    return module


classifier = load_script('classify_wandb_runs')
converter = load_script('convert_csv_to_excel')
tex_updater = load_script('update_tex')
BASELINE = 'Retrieval 미사용'
TARGET = 'target: 16 (anchor 미설정)'


def make_run(run_id, seed, mode, state='running', score=None, hash_bits=10,
             created_at='2026-09-19T00:00:00Z'):
    suffix = '' if mode == 'Both' else ('_O' if mode else '_X')
    return SimpleNamespace(
        id=run_id, name=f'Alien_{run_id}_{seed}{suffix}', state=state,
        commit=None, created_at=created_at,
        summary={} if score is None else {'eval/episode_avg_return': score},
        config={'JointTrainAgent': {'Retrieval': {
            'enable': mode, 'target': 16, 'warmup_steps': 50000,
            'batch_size_reduction': 'retrieved', 'z_score_threshold': 3.5,
            'hash_bits': hash_bits,
        }}},
    )


class RunningRunsTests(unittest.TestCase):
    def setUp(self):
        directory = tempfile.TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        previous_cwd = Path.cwd()
        os.chdir(directory.name)
        self.addCleanup(os.chdir, previous_cwd)

    def export(self, runs):
        with patch.object(classifier.wandb, 'Api') as api, \
                patch.object(classifier.os.path, 'exists', return_value=False), \
                contextlib.redirect_stdout(io.StringIO()):
            api.return_value.runs.return_value = runs
            classifier.main()
        with open('wandb_runs_classification.csv', newline='') as source:
            return {row['Run ID']: row for row in csv.DictReader(source)}

    def workbook_cells(self):
        with contextlib.redirect_stdout(io.StringIO()):
            converter.main(csv_path='wandb_runs_classification.csv',
                           output_path='converted_results.xlsx', queue_dir=Path.cwd())
        workbook = load_workbook('converted_results.xlsx')
        self.addCleanup(workbook.close)
        worksheet = workbook['Results']
        headers = [cell.value for cell in worksheet[1]]
        cells = {}
        game = None
        for row in worksheet.iter_rows(min_row=2):
            if row[0].value:
                game = str(row[0].value).split('\n')[0]
            for header, cell in zip(headers[2:], row[2:]):
                cells[game, row[1].value, header] = cell
        return cells

    def queue_job(self, seed, mode=None, **settings):
        import shlex
        command = [
            'python', '-u', 'train.py', '-n', f'Alien-{seed}', '-seed', str(seed),
            '-env_name', 'ALE/Alien-v5', '-config_path',
            str(converter.HERE.parent / 'config_files/STORM.yaml'),
        ]
        if mode is not None:
            settings['enable'] = repr(mode) if isinstance(mode, list) else str(mode)
        for key, item in settings.items():
            command.extend([f'JointTrainAgent.Retrieval.{key}', str(item)])
        return shlex.join(command)

    def test_queues_include_external_jobs_both_variants_and_new_seeds(self):
        self.export([])
        Path('job_queue.txt').write_text('# manual queue\n\n' + self.queue_job(12345))
        Path('job_queue_custom_gpu.txt').write_text(
            'STORM_AMP_DTYPE=fp16 ' + self.queue_job(12346, ['target1', 'value', 'add']))
        # Stale scheduler history must never turn into a queue marker.
        Path('experiment_scheduler_state.json').write_text(json.dumps({
            'jobs': [{'game': 'Alien', 'seed': 55555, 'status': 'pending'}],
        }))
        cells = self.workbook_cells()
        for config, seed in [(BASELINE, 12345), (TARGET, 12345),
                             ('target: 1 (anchor 미설정)', 12346),
                             (TARGET + ' [value]', 12346), (TARGET + ' [add]', 12346)]:
            cell = cells['Alien', config, seed]
            self.assertEqual(cell.value, 'QUEUED')
            self.assertEqual(cell.fill.fgColor.rgb[-6:], 'DDEBF7')
            self.assertIn('train.py', cell.comment.text)
        self.assertIn('job_queue.txt:3', cells['Alien', TARGET, 12345].comment.text)
        self.assertIsNone(cells['Alien', TARGET, 12346].value)
        self.assertIsNone(cells['Alien', TARGET, converter.PAIRED_MEAN_COLUMN].value)
        self.assertNotIn(('Alien', TARGET, 55555), cells)
        with contextlib.redirect_stdout(io.StringIO()):
            self.assertNotIn('Alien', tex_updater.load_results('converted_results.xlsx'))

    def test_queued_reruns_preserve_scores_hash_preference_and_aggregates(self):
        self.export([
            make_run('baseline', 1, False, 'finished', 100, created_at='2026-08-01T00:00:00Z'),
            make_run('target', 1, True, 'finished', 150, hash_bits=11),
            make_run('baseline2', 2, False, 'finished', 200),
            make_run('target2', 2, True, 'finished', 250),
            make_run('target2older', 2, True, 'finished', 999, hash_bits=11),
        ])
        queue = Path('job_queue_3090.txt')
        commands = [self.queue_job(1, True), self.queue_job(2, True, hash_bits=11)]
        queue.write_text('\n'.join(commands))
        cells = self.workbook_cells()
        self.assertEqual(cells['Alien', TARGET, 1].value, '150.00, QUEUED')
        self.assertEqual(cells['Alien', TARGET, 2].value, '250.00, QUEUED')
        self.assertEqual(cells['Alien', TARGET, converter.PAIRED_MEAN_COLUMN].value, 200)
        self.assertEqual(cells['Alien', BASELINE, converter.PAIRED_MEAN_COLUMN].value, 150)
        self.assertEqual(cells['Alien', TARGET, converter.SCORE_DELTA_COLUMN].value, 50)
        with contextlib.redirect_stdout(io.StringIO()):
            self.assertEqual(tex_updater.load_results('converted_results.xlsx')['Alien'],
                             ([100, 200], [150, 250]))
        self.assertEqual(queue.read_text(), '\n'.join(commands))
        queue.write_text('')
        cells = self.workbook_cells()
        self.assertEqual(cells['Alien', TARGET, 1].value, '150.00')
        self.assertIsNone(cells['Alien', TARGET, 1].fill.fill_type)
        self.assertIsNone(cells['Alien', TARGET, 1].comment)

    def test_queued_styles_preserve_running_exclusions_and_warmup_borders(self):
        self.export([make_run('running', 2, True)])
        command = self.queue_job(5090, True, save_warmup=True)
        Path('job_queue_A6000.txt').write_text('\n'.join([
            self.queue_job(2, True, save_warmup=True), command, command,
        ]))
        cells = self.workbook_cells()
        running = cells['Alien', TARGET, 2]
        self.assertEqual(running.value, 'RUNNING, QUEUED')
        self.assertEqual(running.fill.fgColor.rgb[-6:], 'FFF2CC')
        self.assertEqual(running.border.left.color.rgb[-6:], '00B050')
        excluded = cells['Alien', TARGET, 5090]
        self.assertEqual(excluded.value, 'QUEUED ×2')
        self.assertTrue(excluded.font.strike)
        self.assertEqual(excluded.fill.fgColor.rgb[-6:], 'DDEBF7')
        self.assertEqual(excluded.border.left.color.rgb[-6:], '00B050')
        for note in ('EXCLUDED_SEEDS', 'job_queue_A6000.txt:2', 'job_queue_A6000.txt:3', 'save_warmup'):
            self.assertIn(note, excluded.comment.text)

    def test_queued_configs_use_yaml_and_cli_before_named_overrides(self):
        self.export([])
        config_path = Path('config_files/custom.yaml')
        config_path.parent.mkdir()
        config = converter.yaml.safe_load(
            (converter.HERE.parent / 'config_files/STORM.yaml').read_text())
        config['JointTrainAgent']['Retrieval'].update(target=1, value_signal='value',
                                                     enable=['retrieval', 'add'])
        config_path.write_text(converter.yaml.safe_dump(config))
        Path('job_queue_titan.txt').write_text(
            'python -u train.py -seed=710 -env_name=ALE/Alien-v5 '
            '-config_path=config_files/custom.yaml JointTrainAgent.Retrieval.target 16\n'
            'python train.py -seed 1710 -env_name ALE/Alien-v5 '
            '-config_path config_files/custom.yaml JointTrainAgent.Retrieval.enable "[\'target1\']" '
            'JointTrainAgent.Retrieval.target 16')
        cells = self.workbook_cells()
        self.assertEqual(cells['Alien', TARGET + ' [value]', 710].value, 'QUEUED')
        self.assertEqual(cells['Alien', TARGET + ' [value, add]', 710].value, 'QUEUED')
        self.assertEqual(cells['Alien', 'target: 1 (anchor 미설정) [value]', 1710].value, 'QUEUED')

    def test_queued_resume_reads_saved_config_and_remote_baseline_name(self):
        self.export([])
        warmup = Path('ckpt/arbitrary_name/shared_warmup_50012')
        warmup.mkdir(parents=True)
        (warmup.parent / 'warmup.json').write_text(json.dumps({'checkpoint': warmup.name}))
        (warmup / 'warmup_metadata.json').write_text(json.dumps({
            'run_name': 'arbitrary_name', 'seed': 710, 'env_name': 'ALE/Alien-v5', 'step': 50012,
        }))
        config = converter.yaml.safe_load(
            (converter.HERE.parent / 'config_files/STORM.yaml').read_text())
        config['JointTrainAgent']['Retrieval']['target'] = 1
        (warmup / 'config.yaml').write_text(converter.yaml.safe_dump(config))
        Path('job_queue_pro6k.txt').write_text(
            'python train.py --resume_warmup=ckpt/arbitrary_name '
            'JointTrainAgent.Retrieval.enable "[\'retrieval\', \'baseline\']"\n'
            'python train.py --resume_warmup /remote/STORM/ckpt/Alien-1710_Shared '
            'JointTrainAgent.Retrieval.enable "[\'baseline\']"')
        cells = self.workbook_cells()
        for config, seed in [('target: 1 (anchor 미설정)', 710), (BASELINE, 710), (BASELINE, 1710)]:
            self.assertEqual(cells['Alien', config, seed].value, 'QUEUED')

    def test_remote_resume_marks_all_experiments_without_local_checkpoints(self):
        self.export([])
        paths = [
            'ckpt/Alien-1710',
            'ckpt/Alien-2710_Shared',
            '/remote/STORM/ckpt/Alien-3710/shared_warmup_50012',
            '/remote/STORM/ckpt/Alien-4710_Shared/shared_warmup_50012',
            'ckpt/Alien-seed5710',
        ]
        Path('job_queue_titan.txt').write_text('\n'.join(
            f'python train.py --resume_warmup={path} '
            'JointTrainAgent.Retrieval.enable "[\'baseline\', \'retrieval\', '
            '\'target1\', \'value\', \'add\']"'
            for path in paths))
        with patch.object(converter.warnings, 'warn') as warn:
            cells = self.workbook_cells()
        warn.assert_not_called()
        for seed in (1710, 2710, 3710, 4710, 5710):
            for config in (BASELINE, TARGET, 'target: 1 (anchor 미설정)',
                           TARGET + ' [value]', TARGET + ' [add]'):
                with self.subTest(seed=seed, config=config):
                    self.assertEqual(cells['Alien', config, seed].value, 'QUEUED')
                    self.assertIn('job_queue_titan.txt:', cells['Alien', config, seed].comment.text)

    def test_remote_resume_uses_queue_defaults_and_cli_overrides(self):
        config_path = Path('config_files/STORM.yaml')
        config_path.parent.mkdir()
        config = converter.yaml.safe_load(
            (converter.HERE.parent / 'config_files/STORM.yaml').read_text())
        config['JointTrainAgent']['Retrieval'].update(target=1, value_signal='value')
        config_path.write_text(converter.yaml.safe_dump(config))
        row = converter.queue_command_row(
            'python train.py --resume_warmup ckpt/Alien-1710 '
            'JointTrainAgent.Retrieval.target 16', Path.cwd())
        self.assertEqual(row['Retrieval Enable'], 'Both')
        self.assertEqual(row['Retrieval Target'], 16)
        self.assertEqual(row['Value Signal'], 'value')
        self.assertEqual(row['Warmup Steps'], 50000)
        with self.assertRaisesRegex(ValueError, '설정 파일을 찾을 수 없습니다'):
            converter.queue_command_row(
                'python train.py -n Alien-1710 -config_path missing.yaml', Path.cwd())

    def test_invalid_queue_line_reports_source_and_keeps_other_jobs(self):
        self.export([])
        Path('job_queue_3090.txt').write_text('\n'.join([
            '# comment', 'python train.py "unterminated', self.queue_job(710, ['unknown']),
            self.queue_job(1710, False),
        ]))
        with self.assertWarnsRegex(UserWarning, 'job_queue_3090.txt:2'):
            cells = self.workbook_cells()
        self.assertEqual(cells['Alien', BASELINE, 1710].value, 'QUEUED')
        self.assertIsNone(cells['Alien', BASELINE, 710].value)

    def test_export_keeps_running_and_both_but_excludes_killed_and_finished_both(self):
        both = make_run('both', 5090, 'Both')
        suffix_both = make_run('suffix', 6000, 'Both')
        suffix_both.name += '_Both'
        rows = self.export([
            make_run('running', 710, False), both, suffix_both,
            make_run('killed', 6020, True, state='killed'),
            make_run('doneboth', 2010, 'Both', state='finished', score=999),
        ])
        self.assertEqual(len(rows), 37 + 3)
        self.assertEqual(rows['running']['State'], 'running')
        self.assertEqual(rows['running']['Eval Return'], 'N/A')
        self.assertEqual(rows['both']['Retrieval Enable'], 'Both')
        self.assertEqual(rows['both']['Seed'], '5090')
        self.assertEqual(rows['suffix']['Seed'], '6000')
        self.assertNotIn('killed', rows)
        self.assertNotIn('doneboth', rows)

    def test_workbook_marks_running_and_both_without_using_intermediate_scores(self):
        self.export([
            make_run('base1', 1, False, 'finished', 100, created_at='2026-08-01T00:00:00Z'),
            make_run('target1', 1, True, 'finished', 150),
            make_run('base2', 2, False, 'finished', 200),
            make_run('target2', 2, True, score=9999),
            make_run('base710', 710, False),
            make_run('old710', 710, True, 'finished', 80, created_at='2026-09-18T00:00:00Z'),
            make_run('target710', 710, True, hash_bits=11),
            make_run('both', 5090, 'Both'),
        ])
        cells = self.workbook_cells()
        for config, seed in ((BASELINE, 710), (TARGET, 2), (TARGET, 710),
                             (BASELINE, 5090), (TARGET, 5090)):
            with self.subTest(config=config, seed=seed):
                cell = cells['Alien', config, seed]
                self.assertIn('RUNNING', cell.value)
                self.assertEqual(cell.fill.fgColor.rgb[-6:], 'FFF2CC')
                self.assertEqual(cell.font.color.rgb[-6:], '9C6500')
                self.assertTrue(cell.font.bold)
        self.assertNotIn('9999', cells['Alien', TARGET, 2].value)
        self.assertIn('80.00', cells['Alien', TARGET, 710].value)
        self.assertEqual(cells['Alien', BASELINE, converter.PAIRED_MEAN_COLUMN].value, 100)
        self.assertEqual(cells['Alien', TARGET, converter.PAIRED_MEAN_COLUMN].value, 150)
        self.assertEqual(cells['Alien', BASELINE, 1].value, '100.00')
        self.assertIsNone(cells['Alien', BASELINE, 1].fill.fill_type)

    def test_manual_frostbite_scores_match_scheduler_and_preserve_exclusions(self):
        run = make_run('frost3710', 3710, True, 'finished', 2779)
        run.name = 'Frostbite_frost3710_3710_O'
        self.export([run])
        cells = self.workbook_cells()
        scores = scheduler.latest_scores(
            scheduler.read_rows(Path('wandb_runs_classification.csv')), converter.load_excluded_seeds())
        baseline_score = scores['Frostbite', 'baseline', 3710]['score']
        self.assertEqual(baseline_score, 1904)
        self.assertEqual(converter.parse_score(cells['Frostbite', BASELINE, 3710].value), baseline_score)
        self.assertEqual(cells['Frostbite', BASELINE, converter.PAIRED_MEAN_COLUMN].value, baseline_score)
        self.assertEqual(cells['Frostbite', TARGET, converter.PAIRED_MEAN_COLUMN].value, 2779)
        self.assertEqual(cells['Frostbite', BASELINE, 10].value, '2068.00')
        self.assertTrue(cells['Frostbite', BASELINE, 10].font.strike)
        self.assertNotIn(('Frostbite', 'baseline', 10), scores)

    def test_finished_run_replaces_marker_and_restores_paired_mean(self):
        run = make_run('target', 2, True)
        baseline = make_run('baseline', 2, False, 'finished', 200)
        self.export([baseline, run])
        self.assertEqual(self.workbook_cells()['Alien', TARGET, 2].value, 'RUNNING')
        run.state = 'finished'
        run.summary['eval/episode_avg_return'] = 300
        rows = self.export([baseline, run])
        cells = self.workbook_cells()
        self.assertEqual(rows['target']['State'], 'finished')
        self.assertEqual(cells['Alien', TARGET, 2].value, '300.00')
        self.assertIsNone(cells['Alien', TARGET, 2].fill.fill_type)
        self.assertEqual(cells['Alien', TARGET, converter.PAIRED_MEAN_COLUMN].value, 300)

    def test_variant_configs_export_and_keep_scores_separate(self):
        baseline = make_run('baseline', 2, False, 'finished', 100)
        retrieval = make_run('retrieval', 2, True, 'finished', 150)
        target1 = make_run('target1', 2, True, 'finished', 200)
        target1.config['JointTrainAgent']['Retrieval']['target'] = 1
        value = make_run('value', 2, True, 'finished', 250)
        value.config['JointTrainAgent']['Retrieval'].update(
            value_signal='value', score_combination='multiply', save_warmup=True)
        additive = make_run('add', 2, True, 'finished', 300)
        additive.config['JointTrainAgent.Retrieval.value_signal'] = 'value_diff'
        additive.config['JointTrainAgent.Retrieval.score_combination'] = 'add'
        additive.config['JointTrainAgent.Retrieval.additive_z_score_threshold'] = 4.0
        # add uses its own threshold; the inactive multiply threshold must not filter it.
        additive.config['JointTrainAgent.Retrieval.z_score_threshold'] = 99
        combined = make_run('combined', 2, True, 'finished', 350)
        combined.config['JointTrainAgent']['Retrieval'].update(
            value_signal='value', score_combination='add', additive_z_score_threshold=4.0)
        filtered = make_run('filtered', 2, True, 'finished', 999)
        filtered.config['JointTrainAgent']['Retrieval'].update(
            score_combination='add', additive_z_score_threshold=5.0)

        rows = self.export([baseline, retrieval, target1, value, additive, combined, filtered])
        self.assertEqual(rows['value']['Value Signal'], 'value')
        self.assertEqual(rows['add']['Score Combination'], 'add')
        self.assertEqual(rows['add']['Additive Z Score Threshold'], '4.0')
        cells = self.workbook_cells()
        expected = {
            BASELINE: '100.00', TARGET: '150.00',
            'target: 1 (anchor 미설정)': '200.00',
            TARGET + ' [value]': '250.00', TARGET + ' [add]': '300.00',
            TARGET + ' [value, add]': '350.00',
        }
        for config, score in expected.items():
            self.assertEqual(cells['Alien', config, 2].value, score)
        self.assertEqual(cells['Alien', TARGET, converter.PAIRED_MEAN_COLUMN].value, 150)
        self.assertEqual(cells['Alien', TARGET, converter.SCORE_DELTA_COLUMN].value, 50)
        self.assertEqual(cells['Alien', TARGET + ' [value]', 2].fill.fgColor.rgb[-6:], 'C6EFCE')
        with contextlib.redirect_stdout(io.StringIO()):
            results = tex_updater.load_results('converted_results.xlsx')
        self.assertEqual(results['Alien'], ([100.0], [150.0]))

    def test_tex_does_not_substitute_ablation_scores_for_missing_retrieval(self):
        baseline = make_run('baseline', 2, False, 'finished', 100)
        value = make_run('value', 2, True, 'finished', 900)
        value.config['JointTrainAgent']['Retrieval']['value_signal'] = 'value'
        additive = make_run('add', 2, True, 'finished', 800)
        additive.config['JointTrainAgent']['Retrieval']['score_combination'] = 'add'
        self.export([baseline, value, additive])
        cells = self.workbook_cells()
        self.assertEqual(cells['Alien', TARGET + ' [value]', 2].value, '900.00')
        self.assertEqual(cells['Alien', TARGET + ' [add]', 2].value, '800.00')
        with contextlib.redirect_stdout(io.StringIO()):
            results = tex_updater.load_results('converted_results.xlsx')
        self.assertNotIn('Alien', results)

    def test_shared_experiment_lists_only_mark_selected_effective_configs(self):
        experiments = ['baseline', 'retrieval', 'target1', 'value', 'add']
        shared = make_run('shared', 2, experiments, score=999)
        shared.name = 'Alien_shared_2'
        subset = make_run('subset', 5090, ['value', 'add'])
        inherited = make_run('inherited', 710, ['retrieval', 'add'])
        inherited.config['JointTrainAgent']['Retrieval'].update(target=1, value_signal='value')
        finished = make_run('finished', 6020, experiments, 'finished', 999)
        rows = self.export([shared, subset, inherited, finished])
        self.assertEqual(json.loads(rows['shared']['Retrieval Enable']), experiments)
        self.assertEqual(rows['shared']['Seed'], '2')
        self.assertNotIn('finished', rows)
        cells = self.workbook_cells()
        for config in [BASELINE, TARGET, 'target: 1 (anchor 미설정)',
                       TARGET + ' [value]', TARGET + ' [add]']:
            cell = cells['Alien', config, 2]
            self.assertEqual(cell.value, 'RUNNING')
            self.assertEqual(cell.fill.fgColor.rgb[-6:], 'FFF2CC')
        for config in [TARGET + ' [value]', TARGET + ' [add]']:
            self.assertEqual(cells['Alien', config, 5090].value, 'RUNNING')
        self.assertIsNone(cells['Alien', BASELINE, 5090].value)
        self.assertIsNone(cells['Alien', TARGET, 5090].value)
        self.assertEqual(cells['Alien', 'target: 1 (anchor 미설정) [value]', 710].value, 'RUNNING')
        self.assertEqual(cells['Alien', 'target: 1 (anchor 미설정) [value, add]', 710].value, 'RUNNING')
        self.assertTrue(all('999' not in str(cell.value) for cell in cells.values()))

    def test_running_branch_marks_only_current_and_remaining_experiments(self):
        experiments = ['retrieval', 'target1', 'value', 'add']
        retrieval = make_run('retrieval', 6030, True, 'finished', 100)
        target1 = make_run('target1', 6030, True, 'finished', 200)
        target1.config['JointTrainAgent']['Retrieval']['target'] = 1
        value = make_run('value', 6030, True, score=9999)
        value.config['JointTrainAgent']['Retrieval'].update(value_signal='value', save_warmup=False)
        value.file = Mock()
        args = ['JointTrainAgent.Retrieval.enable', repr(experiments),
                'JointTrainAgent.Retrieval.save_warmup', 'True',
                '--branch_experiment', 'value']
        value.file.return_value.download.side_effect = lambda **kwargs: io.StringIO(
            json.dumps({'args': args}))

        rows = self.export([retrieval, target1, value])
        self.assertEqual(json.loads(rows['value']['Pending Retrieval Configs']), [
            {'Retrieval Enable': True, 'Value Signal': 'value_diff', 'Score Combination': 'add'},
        ])
        value.file.assert_called_once_with('wandb-metadata.json')
        cells = self.workbook_cells()
        self.assertEqual(cells['Alien', TARGET, 6030].value, '100.00')
        self.assertEqual(cells['Alien', 'target: 1 (anchor 미설정)', 6030].value, '200.00')
        for config in [TARGET + ' [value]', TARGET + ' [add]']:
            cell = cells['Alien', config, 6030]
            self.assertEqual(cell.value, 'RUNNING')
            self.assertEqual(cell.fill.fgColor.rgb[-6:], 'FFF2CC')
            self.assertEqual(cell.border.left.color.rgb[-6:], '00B050')
        self.assertNotIn(('Alien', TARGET + ' [value, add]', 6030), cells)
        self.assertTrue(all('9999' not in str(cell.value) for cell in cells.values()))

        value.state = 'finished'
        value.summary['eval/episode_avg_return'] = 250
        additive = make_run('add', 6030, True)
        additive.config['JointTrainAgent']['Retrieval']['score_combination'] = 'add'
        additive.file = Mock()
        additive.file.return_value.download.return_value = io.StringIO(json.dumps({
            'args': args[:-1] + ['add'],
        }))
        rows = self.export([retrieval, target1, value, additive])
        self.assertEqual(json.loads(rows['value']['Pending Retrieval Configs']), [])
        self.assertEqual(json.loads(rows['add']['Pending Retrieval Configs']), [])
        cells = self.workbook_cells()
        self.assertEqual(cells['Alien', TARGET + ' [value]', 6030].value, '250.00')
        self.assertEqual(cells['Alien', TARGET + ' [add]', 6030].value, 'RUNNING')

        additive.state = 'finished'
        additive.summary['eval/episode_avg_return'] = 300
        self.export([retrieval, target1, value, additive])
        self.assertEqual(self.workbook_cells()['Alien', TARGET + ' [add]', 6030].value, '300.00')

    def test_pending_branches_follow_list_order_and_restore_common_cli_settings(self):
        for branch, settings, extra_args, expected in [
            ('target1', {'target': 1}, [], [
                TARGET, TARGET + ' [value]', TARGET + ' [add]', BASELINE,
            ]),
            ('value', {'target': 1, 'value_signal': 'value'}, [
                'JointTrainAgent.Retrieval.value_signal', 'value_diff',
                'JointTrainAgent.Retrieval.value_signal', 'value',
            ], ['target: 1 (anchor 미설정) [value, add]', BASELINE]),
            ('add', {'score_combination': 'add'}, [], [BASELINE]),
        ]:
            with self.subTest(branch=branch):
                run = make_run('child', 6030, True)
                run.config['JointTrainAgent']['Retrieval'].update(settings)
                run.file = Mock()
                run.file.return_value.download.return_value = io.StringIO(json.dumps({'args': [
                    'JointTrainAgent.Retrieval.enable', repr(['target1', 'retrieval', 'value', 'add', 'baseline']),
                    *extra_args, f'--branch_experiment={branch}',
                ]}))
                self.export([run])
                cells = self.workbook_cells()
                for config in expected:
                    self.assertEqual(cells['Alien', config, 6030].value, 'RUNNING')

    def test_pending_configs_ignore_missing_invalid_and_unrelated_metadata(self):
        for args in (None, [], ['--branch_experiment', 'value'],
                     ['JointTrainAgent.Retrieval.enable', 'invalid', '--branch_experiment', 'value'],
                     ['JointTrainAgent.Retrieval.enable', "['add']", '--branch_experiment', 'value'],
                     ['JointTrainAgent.Retrieval.enable', "['value', 'unknown']", '--branch_experiment', 'value']):
            with self.subTest(args=args):
                self.assertEqual(classifier.parse_pending_retrieval_configs(args), [])

    def test_child_warmup_request_highlights_all_five_effective_configs(self):
        # Real failure: koenosnj requested save_warmup=True on the CLI, but
        # configure_storm_retrieval_run resets the child's logged config to False.
        variants = {
            'baseline': (False, {}, BASELINE),
            'retrieval': (True, {}, TARGET),
            'target1': (True, {'target': 1}, 'target: 1 (anchor 미설정)'),
            'value': (True, {'value_signal': 'value'}, TARGET + ' [value]'),
            'add': (True, {'score_combination': 'add'}, TARGET + ' [add]'),
        }
        runs = []
        for experiment, (enabled, overrides, _) in variants.items():
            run_id = 'koenosnj' if experiment == 'retrieval' else experiment
            run = make_run(run_id, 6030, enabled, 'finished', 2933)
            run.name = f'Gopher_{run_id}_6030'
            run.config['JointTrainAgent']['Retrieval'].update(save_warmup=False, **overrides)
            run.file = Mock()
            run.file.return_value.download.return_value = io.StringIO(json.dumps({'args': [
                '-n', 'Gopher-6030', '-seed', '6030',
                'JointTrainAgent.Retrieval.enable', repr([experiment]),
                'JointTrainAgent.Retrieval.save_warmup', 'True',
                '--branch_experiment', experiment,
                '--resume_from', '/home/choemj/STORM-1/ckpt/Gopher-6030_Shared/shared_warmup_50356',
            ]}))
            runs.append(run)

        rows = self.export(runs)
        # This fixture tests warmup highlighting independently of the user's
        # editable exclusion list (Gopher 6030 may now legitimately be excluded).
        with patch.object(converter, 'load_excluded_seeds', return_value={}):
            cells = self.workbook_cells()
        for run, (_, _, config) in zip(runs, variants.values()):
            with self.subTest(config=config):
                self.assertEqual(rows[run.id]['Save Warmup'], 'False')
                self.assertEqual(rows[run.id]['Save Warmup Requested'], 'True')
                self.assertEqual(rows[run.id]['Warmup Directory'],
                                 '/home/choemj/STORM-1/ckpt/Gopher-6030_Shared/shared_warmup_50356')
                self.assertEqual(rows[run.id]['Base Run Name'], 'Gopher-6030')
                self.assertEqual(rows[run.id]['Training Phase'], 'branch')
                cell = cells['Gopher', config, 6030]
                self.assertEqual(cell.value, '2933.00')
                self.assertEqual(cell.fill.fgColor.rgb[-6:], 'C6EFCE')
                self.assertEqual(cell.border.left.color.rgb[-6:], '00B050')

    def test_missing_or_false_warmup_request_keeps_unhighlighted_score(self):
        key = 'JointTrainAgent.Retrieval.save_warmup'
        runs = []
        for seed, args in enumerate((None, [], [key, 'True', key, 'False']), 6030):
            run = make_run(str(seed), seed, True, 'finished', 2933)
            run.config['JointTrainAgent']['Retrieval']['save_warmup'] = False
            run.file = Mock()
            if args is None:
                run.file.side_effect = FileNotFoundError('No metadata')
            else:
                run.file.return_value.download.return_value = io.StringIO(json.dumps({'args': args}))
            runs.append(run)
        rows = self.export(runs)
        cells = self.workbook_cells()
        self.assertEqual(rows['6030']['Save Warmup Requested'], 'N/A')
        self.assertEqual(rows['6031']['Save Warmup Requested'], 'N/A')
        self.assertEqual(rows['6032']['Save Warmup Requested'], 'False')
        for seed in range(6030, 6033):
            cell = cells['Alien', TARGET, seed]
            self.assertEqual(cell.value, '2933.00')
            self.assertIsNone(cell.fill.fill_type)

    def test_running_child_preserves_yellow_and_adds_warmup_border(self):
        run = make_run('child', 6030, True, score=999)
        run.config['JointTrainAgent']['Retrieval']['save_warmup'] = False
        run.file = Mock()
        run.file.return_value.download.return_value = io.StringIO(json.dumps({'args': [
            'JointTrainAgent.Retrieval.save_warmup', 'True',
        ]}))
        self.export([run])
        cell = self.workbook_cells()['Alien', TARGET, 6030]
        self.assertEqual(cell.value, 'RUNNING')
        self.assertEqual(cell.fill.fgColor.rgb[-6:], 'FFF2CC')
        self.assertEqual(cell.border.left.color.rgb[-6:], '00B050')


if __name__ == '__main__':
    unittest.main()
