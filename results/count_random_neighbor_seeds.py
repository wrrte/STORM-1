"""RandomNeighor 결과와 Full 결과의 게임별 시드 커버리지를 Excel로 저장합니다.

실행 예시 (STORM-1 디렉터리에서):
    python results/count_random_neighbor_seeds.py
    python results/count_random_neighbor_seeds.py --input results/converted_results.xlsx \
        --output results/random_neighbor_seed_coverage.xlsx

의존성: openpyxl (python -m pip install openpyxl)
원본 Excel은 수정하지 않습니다. 기본 경로는 현재 작업 폴더와 무관합니다.
"""

import argparse
from dataclasses import dataclass
import math
from pathlib import Path
import re

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


FULL_CONFIG = "target: 16 (anchor 미설정)"
RANDOM_CONFIG = FULL_CONFIG + " [RandomNeighor]"
HERE = Path(__file__).resolve().parent
NUMBER = re.compile(r"[+-]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?")


@dataclass(frozen=True)
class SeedResult:
    text: str = ""
    scores: tuple = ()
    excluded: bool = False
    coordinate: str = ""

    @property
    def running(self):
        return "RUNNING" in self.text.upper()

    @property
    def queued(self):
        return "QUEUED" in self.text.upper()


def parse_scores(value):
    """괄호 안 hb/w 메타데이터를 제거한 뒤 완료 점수만 추출합니다.

    '10.0 (hb: 10, w: 50000), 11.0, QUEUED' -> (10.0, 11.0).
    RUNNING/QUEUED와 숫자가 함께 있으면 완료 점수는 계속 인정합니다.
    """
    if value is None or isinstance(value, bool):
        return ()
    text = re.sub(r"\([^()]*\)", "", str(value))
    scores = []
    for item in text.split(","):
        item = item.strip()
        if NUMBER.fullmatch(item):
            score = float(item)
            if math.isfinite(score):
                scores.append(score)
    return tuple(scores)


def seed_id(value):
    """시드 헤더만 인식하며 공백 구분 열과 평균/차이 열은 무시합니다."""
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float)):
        if math.isfinite(value) and value >= 0 and int(value) == value:
            return int(value)
    elif str(value).strip().isdigit():
        return int(str(value).strip())
    return None


def read_results(input_path, sheet_name="Results"):
    """병합된 Game 셀을 이어 읽고 정확히 일치하는 두 설정만 수집합니다."""
    workbook = load_workbook(input_path, data_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"시트가 없습니다: {sheet_name}; 사용 가능: {workbook.sheetnames}")
        sheet = workbook[sheet_name]
        headers = {cell.value: cell.column for cell in sheet[1] if cell.value is not None}
        if not {"Game", "Config"} <= headers.keys():
            raise ValueError("첫 행에 Game, Config 헤더가 필요합니다.")
        seed_columns = []
        seen_seeds = set()
        for cell in sheet[1]:
            seed = seed_id(cell.value)
            if seed is not None:
                if seed in seen_seeds:
                    raise ValueError(f"중복된 시드 헤더: {seed}")
                seed_columns.append((cell.column, seed))
                seen_seeds.add(seed)
        if not seed_columns:
            raise ValueError("숫자로 된 시드 헤더가 없습니다.")

        games = {}
        game = None
        for row in sheet.iter_rows(min_row=2):
            game_value = row[headers["Game"] - 1].value
            if game_value is not None and str(game_value).strip():
                # 원본에는 게임명 아래 Random/Human 참조 점수가 들어 있습니다.
                game = str(game_value).strip().splitlines()[0].strip()
            if game is None:
                continue
            methods = games.setdefault(game, {})
            config = row[headers["Config"] - 1].value
            if config not in {FULL_CONFIG, RANDOM_CONFIG}:
                continue
            if config in methods:
                raise ValueError(f"같은 게임/설정의 행이 중복되었습니다: {game}, {config}")
            methods[config] = {
                seed: SeedResult(
                    text="" if row[column - 1].value is None else str(row[column - 1].value),
                    scores=parse_scores(row[column - 1].value),
                    excluded=bool(row[column - 1].font.strike),
                    coordinate=row[column - 1].coordinate,
                )
                for column, seed in seed_columns
            }
        if not games:
            raise ValueError("게임 행이 없습니다.")
        if not any(methods for methods in games.values()):
            raise ValueError("Full 또는 RandomNeighor 설정 행이 없습니다. 설정명을 확인하세요.")
        return games
    finally:
        workbook.close()


def seed_text(seeds):
    return ", ".join(str(seed) for seed in sorted(seeds)) or "-"


def goal_status(count, minimum, maximum):
    if count < minimum:
        return "최소 미달"
    if count > maximum:
        return "최대 초과"
    return "목표 범위"


def build_tables(games, minimum=2, maximum=4):
    summary = []
    details = []
    empty = SeedResult()
    for game, methods in games.items():
        full = methods.get(FULL_CONFIG, {})
        random = methods.get(RANDOM_CONFIG, {})
        full_seeds = {seed for seed, result in full.items() if result.scores}
        random_seeds = {seed for seed, result in random.items() if result.scores}
        paired = full_seeds & random_seeds
        excluded = {seed for seed in full.keys() | random.keys()
                    if full.get(seed, empty).excluded or random.get(seed, empty).excluded}
        usable_paired = paired - excluded
        summary.append([
            game, len(random_seeds), len(paired), len(random_seeds - full_seeds),
            goal_status(len(paired), minimum, maximum), max(0, minimum - len(paired)),
            max(0, maximum - len(paired)), len(usable_paired),
            goal_status(len(usable_paired), minimum, maximum),
            sum(len(result.scores) for result in random.values()),
            seed_text(random_seeds), seed_text(paired), seed_text(random_seeds - full_seeds),
            seed_text(full_seeds - random_seeds - excluded),
            seed_text(seed for seed, result in random.items() if result.running),
            seed_text(seed for seed, result in random.items() if result.queued),
            seed_text(paired & excluded),
        ])

        # Full만 있는 시드도 표시하여 다음 ablation 실행 후보를 확인할 수 있습니다.
        for seed in sorted(full.keys() | random.keys()):
            f, r = full.get(seed, empty), random.get(seed, empty)
            if not f.text.strip() and not r.text.strip():
                continue
            if r.scores and f.scores:
                status = "대응 완료"
            elif r.scores:
                status = "Full 완료 결과 없음"
            elif f.scores:
                status = "RandomNeighor 완료 결과 없음"
            else:
                status = "완료 결과 없음"
            details.append([
                game, seed, "있음" if r.scores else "없음", "있음" if f.scores else "없음",
                status, "예" if seed in excluded else "아니오",
                "예" if seed in usable_paired else "아니오", len(r.scores),
                r.text, f.text, "예" if r.running else "아니오",
                "예" if r.queued else "아니오", r.coordinate, f.coordinate,
            ])
    return summary, details


def write_table(workbook, title, headers, rows, widths):
    sheet = workbook.create_sheet(title)
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    sheet.freeze_panes = "B2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[1].height = 48
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    colors = {"최소 미달": "FFF2CC", "최대 초과": "F8CBAD", "목표 범위": "E2F0D9"}
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if isinstance(cell.value, str) and cell.value in colors:
                cell.fill = PatternFill("solid", fgColor=colors[cell.value])
    return sheet


def write_report(input_path, output_path, games, minimum=2, maximum=4, sheet_name="Results"):
    if not 1 <= minimum <= maximum:
        raise ValueError("1 <= min-seeds <= max-seeds 조건을 만족해야 합니다.")
    if Path(input_path).resolve() == Path(output_path).resolve():
        raise ValueError("원본과 출력 파일 경로가 같습니다. 다른 출력 경로를 지정하세요.")
    summary, details = build_tables(games, minimum, maximum)
    workbook = Workbook()
    workbook.remove(workbook.active)
    summary_sheet = write_table(workbook, "게임별 요약", [
        "Game", "RandomNeighor 결과 시드 수", "Full과 공통 결과 시드 수",
        "Full 미대응 결과 시드 수", "목표 상태 (공통 시드)",
        f"최소 {minimum}개까지 부족", f"{maximum}개까지 추가 가능",
        "공통 시드 수 (취소선 제외)", "목표 상태 (취소선 제외)",
        "RandomNeighor 완료 run 수 (중복 포함)", "RandomNeighor 완료 시드",
        "Full과 공통 완료 시드", "Full 없는 RandomNeighor 시드",
        "RandomNeighor 없는 Full 시드 (취소선 제외)", "RandomNeighor RUNNING 시드",
        "RandomNeighor QUEUED 시드", "공통 시드 중 취소선 시드",
    ], summary, [22, 23, 23, 23, 19, 18, 18, 23, 22, 25, 35, 30, 35, 45, 30, 30, 30])
    for column in ("B1", "C1", "D1"):
        summary_sheet[column].comment = Comment(
            "완료된 숫자 점수가 하나 이상 있는 고유 시드 수입니다. "
            "동일 시드의 중복 run은 1개로 셉니다. 취소선 시드도 포함합니다.", "STORM"
        )
    write_table(workbook, "시드별 상세", [
        "Game", "Seed", "RandomNeighor 완료 결과", "Full 완료 결과", "대응 상태",
        "취소선 제외 시드", "취소선 제외 후 대응 완료", "RandomNeighor 완료 run 수",
        "RandomNeighor 원본 셀", "Full 원본 셀", "RandomNeighor RUNNING",
        "RandomNeighor QUEUED", "RandomNeighor 원본 셀 위치", "Full 원본 셀 위치",
    ], details, [22, 12, 24, 20, 38, 20, 25, 25, 55, 55, 25, 25, 26, 24])
    guide = workbook.create_sheet("집계 기준")
    for row in [
        ("항목", "설명"),
        ("입력 파일", str(Path(input_path).resolve())),
        ("입력 시트", sheet_name),
        ("Full 설정 (정확히 일치)", FULL_CONFIG),
        ("Ablation 설정 (정확히 일치)", RANDOM_CONFIG),
        ("완료 결과 시드", "숫자 점수가 하나 이상 있는 고유 시드. 음수와 0도 유효합니다."),
        ("중복 실행", "같은 시드의 여러 점수는 시드 수로는 1개, 완료 run 수로는 각각 셉니다."),
        ("점수 파싱", "괄호 안 hb/w 메타데이터를 무시하고 쉼표로 구분된 유한 숫자만 셉니다."),
        ("미완료 셀", "빈 셀, N/A, NaN, RUNNING, QUEUED만 있는 셀은 완료 결과에서 제외합니다."),
        ("완료/미완료 혼합 셀", "숫자와 RUNNING/QUEUED가 함께 있으면 완료 결과와 진행 상태를 모두 표시합니다."),
        ("공통 시드", "같은 게임, 같은 시드에서 Full과 RandomNeighor 모두 완료 점수가 있는 경우입니다."),
        ("목표와 부족/추가 수", f"전체 공통 완료 시드 수를 기준으로 최소 {minimum}, 최대 {maximum}개를 확인합니다. 초과 결과도 모두 표시합니다."),
        ("취소선 제외 집계", "원본의 두 설정 중 어느 셀이든 취소선이면 제외합니다. 원본 Excel의 제외 표기를 그대로 따릅니다."),
        ("추가 실험 후보", "RandomNeighor 없는 Full 시드는 완료된 Full 점수가 있으나 ablation 완료 점수가 없는 비제외 시드입니다. 진행 중인 시드도 포함될 수 있습니다."),
        ("설정 행이 없는 게임", "다른 설정 행에 게임이 있으면 포함하며, 없는 설정의 결과 수는 0입니다."),
        ("시드별 상세", "두 설정 중 하나라도 셀 값이 있는 시드만 표시합니다. 원본 셀과 좌표로 확인할 수 있습니다."),
        ("결과 보존", "원본을 수정하지 않으며, 최대 시드 수를 넘어도 잘라내거나 점수로 시드를 고르지 않습니다."),
    ]:
        guide.append(row)
    guide.column_dimensions["A"].width = 30
    guide.column_dimensions["B"].width = 120
    guide.freeze_panes = "A2"
    for row in guide:
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    workbook.close()
    return summary


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--input", type=Path, default=HERE / "converted_results.xlsx")
    parser.add_argument("--output", type=Path, default=HERE / "random_neighbor_seed_coverage.xlsx")
    parser.add_argument("--sheet", default="Results")
    parser.add_argument("--min-seeds", type=int, default=2)
    parser.add_argument("--max-seeds", type=int, default=4)
    args = parser.parse_args()
    if not 1 <= args.min_seeds <= args.max_seeds:
        parser.error("1 <= min-seeds <= max-seeds 조건을 만족해야 합니다.")
    games = read_results(args.input, args.sheet)
    summary = write_report(args.input, args.output, games, args.min_seeds, args.max_seeds, args.sheet)
    print("Game                  Random 시드  공통 시드  공통(취소선 제외)")
    for row in summary:
        print(f"{row[0]:22s} {row[1]:10d} {row[2]:10d} {row[7]:18d}")
    reached = sum(args.min_seeds <= row[2] <= args.max_seeds for row in summary)
    print(f"목표 범위에 있는 게임: {reached}/{len(summary)}")
    print(f"저장: {args.output.resolve()}")


if __name__ == "__main__":
    main()
