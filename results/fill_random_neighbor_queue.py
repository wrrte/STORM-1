"""2000~3999번 시드의 RandomNeighor 실험을 게임당 4개까지 보충합니다.

STORM-1에서 실행:
    python results/fill_random_neighbor_queue.py          # 계획만 확인
    python results/fill_random_neighbor_queue.py --apply  # 실제 큐에 추가

각 게임의 3개째 명령을 모두 배치한 뒤 4개째 명령을 배치합니다. 3개보다
적게 확보한 게임은 앞선 차수부터 보충합니다. Krull은 항상 제외하고,
Kangaroo의 모든 추가 명령은 다른 게임의 4개째 명령 뒤에 배치합니다.
완료/실행 중/대기 중인 동일 게임·시드는 중복 등록하지 않습니다.
Full 결과가 있는 후보는 논문과의 HNS 차이로 정렬하고,
후보가 부족하면 같은 게임에서 아직 사용하지 않은 시드를 배정합니다.
취소선 시드는 새 실험 후보에서 제외합니다. Full 중복 점수는 원본 셀의
첫 번째 완료 점수(최신 순)를 사용합니다. 시드 수는 중복 run을 합쳐 셉니다.
Freeway는 trajectory 경로를 Freeway.pkl로 지정하고 UseDemonstration을 켭니다.

기존 Excel 변환기의 큐/YAML 파서를 재사용하므로 openpyxl, pandas, PyYAML이
필요합니다. 학습 명령은 기존 worker가 큐에서 가져가 실행합니다.
"""

import argparse
import ast
from collections import defaultdict
from contextlib import ExitStack, contextmanager
import fcntl
from itertools import chain
import json
import math
import os
from pathlib import Path
import re
import shlex

from openpyxl import load_workbook

from count_random_neighbor_seeds import FULL_CONFIG, RANDOM_CONFIG, read_results, seed_id
from convert_csv_to_excel import config_value, queue_command_row


HERE = Path(__file__).resolve().parent
SEED_MIN, SEED_MAX = 2000, 3999
NUMBER = r"[+-]?(?:\d[\d,]*(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?"
PREFIX = "JointTrainAgent.Retrieval."


def in_range(seed):
    return SEED_MIN <= seed <= SEED_MAX


def workbook_metadata(input_path, sheet_name):
    """논문/Random/Human 점수와 모든 설정에서 사용한 시드를 원본에서 읽습니다."""
    references, used = {}, defaultdict(set)
    workbook = load_workbook(input_path, data_only=True)
    try:
        sheet = workbook[sheet_name]
        columns = {cell.value: cell.column for cell in sheet[1] if cell.value is not None}
        seeds = [(cell.column, seed_id(cell.value)) for cell in sheet[1]
                 if seed_id(cell.value) is not None]
        game = None
        for row in sheet.iter_rows(min_row=2):
            label = row[columns["Game"] - 1].value
            if label is not None and str(label).strip():
                game = str(label).strip().splitlines()[0].strip()
                reference = references.setdefault(game, {})
                for key in ("Random", "Human"):
                    match = re.search(rf"(?:^|\n)\s*{key}:\s*({NUMBER})\s*(?:\n|$)", str(label))
                    if match:
                        reference[key.lower()] = float(match.group(1).replace(",", ""))
            if game is None:
                continue
            config = row[columns["Config"] - 1].value
            match = re.fullmatch(rf"STORM \(논문\):\s*({NUMBER})\s*", str(config))
            if match:
                references.setdefault(game, {})["paper"] = float(match.group(1).replace(",", ""))
            for column, seed in seeds:
                if row[column - 1].value is not None and str(row[column - 1].value).strip():
                    used[game].add(seed)
    finally:
        workbook.close()
    return references, used


def is_target_random(row):
    """YAML/CLI를 반영한 큐 설정이 정확한 RandomNeighor 행인지 판정합니다."""
    mode = str(row["Retrieval Enable"]).strip().lower()
    if mode.startswith("["):
        names = ast.literal_eval(mode)
        if "randomneighor" not in names:
            return False
        row = {**row, "Neighbor Source": "random"}
    elif mode == "randomneighor":
        row = {**row, "Neighbor Source": "random"}
    elif mode not in {"true", "both"}:
        return False
    expected = {
        "Warmup Steps": 50000, "Retrieval Target": 16,
        "Z Score Threshold": 3.5, "Anchor Weight": 0.0625,
    }
    for key, value in expected.items():
        try:
            if float(config_value(row, key, value)) != value:
                return False
        except (TypeError, ValueError):
            return False
    defaults = {"Neighbor Source": "retrieval"}
    return all(str(config_value(row, key, defaults.get(key, expected))).strip().lower() == expected
               for key, expected in {
        "Batch Size Reduction": "retrieved", "Value Signal": "value_diff",
        "Score Combination": "multiply", "Neighbor Source": "random",
    }.items())


def queue_reservations(queue_texts):
    """모든 현재 큐의 시드를 예약합니다. 명령을 실행하지 않습니다."""
    random_pending, used = defaultdict(set), defaultdict(set)
    for path, text in queue_texts.items():
        for number, line in enumerate(text.splitlines(), 1):
            try:
                row = queue_command_row(line, Path(path).parent)
            except Exception as error:
                raise ValueError(f"{path}:{number}: 큐 명령을 읽을 수 없습니다: {error}") from error
            if row is None:
                continue
            game = row["Run Name"].rsplit("_queued_", 1)[0]
            seed = int(row["Seed"])
            used[game].add(seed)
            if is_target_random(row):
                random_pending[game].add(seed)
    return random_pending, used


def make_command(game, seed):
    trajectory = "Freeway.pkl" if game == "Freeway" else f"D_TRAJ/{game}.pkl"
    args = [
        "python", "-u", "train.py", "-n", f"{game}-{seed}", "-seed", str(seed),
        "-config_path", "config_files/STORM.yaml", "-env_name", f"ALE/{game}-v5",
        "-trajectory_path", trajectory,
    ]
    if game == "Freeway":
        args.extend(["JointTrainAgent.UseDemonstration", "True"])
    # YAML 기본값이 바뀌어도 요청한 설정 행으로 분류되도록 고정합니다.
    for key, value in {
        "enable": "RandomNeighor", "target": "16", "anchor_weight": "0.0625",
        "warmup_steps": "50000", "batch_size_reduction": "retrieved",
        "z_score_threshold": "3.5", "value_signal": "value_diff",
        "score_combination": "multiply", "neighbor_source": "random",
        "hash_bits": "10", "save_warmup": "False",
    }.items():
        args.extend([PREFIX + key, value])
    return shlex.join(args)


def plan_jobs(games, references, workbook_used, pending, queue_used, target=4):
    if not 1 <= target <= 4:
        raise ValueError("target-seeds는 1~4여야 합니다.")
    jobs, summaries = [], []
    for game, methods in games.items():
        if game == "Krull":
            continue
        full = methods.get(FULL_CONFIG, {})
        random = methods.get(RANDOM_CONFIG, {})
        completed = {seed for seed, result in random.items() if in_range(seed) and result.scores}
        reserved = {seed for seed, result in random.items()
                    if in_range(seed) and (result.running or result.queued)}
        reserved |= {seed for seed in pending.get(game, set()) if in_range(seed)}
        covered = completed | reserved
        needed = max(0, target - len(covered))
        if not needed:
            continue

        excluded = {seed for results in (full, random) for seed, result in results.items()
                    if result.excluded}
        unavailable = covered | excluded | {seed for seed, result in random.items() if result.text.strip()}
        candidates = [seed for seed, result in full.items()
                      if in_range(seed) and result.scores and seed not in unavailable]
        reference = references.get(game, {})
        if candidates:
            if not {"random", "human", "paper"} <= reference.keys():
                raise ValueError(f"{game}: Full 후보의 HNS 비교에 필요한 논문/Random/Human 점수가 없습니다.")
            if not all(math.isfinite(value) for value in reference.values()):
                raise ValueError(f"{game}: 참조 점수가 유한수가 아닙니다.")
            denominator = reference["human"] - reference["random"]
            if denominator == 0:
                raise ValueError(f"{game}: Human과 Random 점수가 같아 HNS를 계산할 수 없습니다.")
            candidates.sort(key=lambda seed: (
                abs((full[seed].scores[0] - reference["paper"]) / denominator), seed,
            ))
        selected = candidates[:needed]
        # 후보 Full 시드는 재사용하되, fallback은 다른 설정까지 포함해 미사용 시드만 배정합니다.
        used = set(workbook_used.get(game, set())) | set(queue_used.get(game, set())) | unavailable
        for _ in range(needed - len(selected)):
            available = chain(range(SEED_MIN, SEED_MAX + 1, 10),
                              (seed for seed in range(SEED_MIN, SEED_MAX + 1) if seed % 10))
            seed = next((seed for seed in available if seed not in used and seed not in selected), None)
            if seed is None:
                raise ValueError(f"{game}: {SEED_MIN}~{SEED_MAX}에 사용 가능한 시드가 없습니다.")
            selected.append(seed)
            used.add(seed)
        summaries.append({"game": game, "completed": sorted(completed),
                          "pending": sorted(reserved - completed), "new_seeds": selected,
                          "expected_seed_count": len(covered) + len(selected)})
        for round_target, seed in enumerate(selected, start=len(covered) + 1):
            matched = seed in candidates
            full_score = full[seed].scores[0] if matched else None
            paper_hns = ((reference["paper"] - reference["random"]) /
                         (reference["human"] - reference["random"])) if matched else None
            full_hns = ((full_score - reference["random"]) /
                        (reference["human"] - reference["random"])) if matched else None
            jobs.append({
                "game": game, "seed": seed, "round_target": round_target,
                "reason": "Full HNS가 논문에 가까운 순" if matched else "Full 후보가 없어 미사용 시드 배정",
                "full_score": full_score, "paper_score": reference.get("paper"),
                "full_hns": full_hns, "paper_hns": paper_hns,
                "absolute_hns_gap": abs(full_hns - paper_hns) if matched else None,
                "command": make_command(game, seed),
            })
    # 모든 게임의 n개째 명령을 배치한 뒤 n+1개째 명령을 배치합니다.
    # Kangaroo는 모든 차수를 맨 뒤로 옮기고, 같은 차수 안에서는 Excel 순서를 유지합니다.
    jobs.sort(key=lambda job: (job["game"] == "Kangaroo", job["round_target"]))
    summaries.sort(key=lambda summary: summary["game"] == "Kangaroo")
    return {"seed_range": [SEED_MIN, SEED_MAX], "target_seeds": target,
            "excluded_games": ["Krull"], "games": summaries, "new_jobs": jobs}


@contextmanager
def locked_queues(queue_path):
    """기존 GPU worker와 동일한 flock 파일로 큐 읽기/추가를 보호합니다."""
    queue_path = Path(queue_path).resolve()
    paths = set(queue_path.parent.glob("job_queue_*.txt")) | {queue_path}
    default = queue_path.parent / "job_queue.txt"
    if default.exists():
        paths.add(default)
    with ExitStack() as stack:
        for path in sorted(paths):
            lock = stack.enter_context(path.with_suffix(".lock").open("a"))
            fcntl.flock(lock, fcntl.LOCK_EX)
        texts = {path: path.read_text(encoding="utf-8") if path.exists() else "" for path in paths}
        yield texts


def run(input_path, queue_path, sheet_name="Results", target=4, apply=False):
    queue_path = Path(queue_path).resolve()
    if queue_path == Path(input_path).resolve():
        raise ValueError("입력 Excel과 큐는 다른 경로여야 합니다.")
    with locked_queues(queue_path) as texts:
        games = read_results(input_path, sheet_name)
        references, used = workbook_metadata(input_path, sheet_name)
        pending, queue_used = queue_reservations(texts)
        report = plan_jobs(games, references, used, pending, queue_used, target)
        commands = [job["command"] for job in report["new_jobs"]]
        if apply and commands:
            with queue_path.open("a", encoding="utf-8") as output:
                if texts[queue_path] and not texts[queue_path].endswith("\n"):
                    output.write("\n")
                output.write("\n".join(commands) + "\n")
                output.flush()
                os.fsync(output.fileno())
    report.update({"input": str(Path(input_path).resolve()), "queue": str(queue_path), "applied": apply})
    return report


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--input", type=Path, default=HERE / "converted_results.xlsx")
    parser.add_argument("--queue", type=Path, default=HERE.parent / "job_queue_3090.txt")
    parser.add_argument("--sheet", default="Results")
    parser.add_argument("--target-seeds", type=int, choices=range(1, 5), default=4)
    parser.add_argument("--apply", action="store_true", help="실제로 큐에 추가 (기본은 계획만 표시)")
    parser.add_argument("--report", type=Path, help="선택 근거와 명령을 JSON으로 저장")
    args = parser.parse_args()
    if args.report and args.report.resolve() in {args.input.resolve(), args.queue.resolve()}:
        parser.error("report는 입력 Excel/큐와 다른 경로여야 합니다.")
    report = run(args.input, args.queue, args.sheet, args.target_seeds, args.apply)
    for game in report["games"]:
        print(f"{game['game']}: 완료 {len(game['completed'])}, 대기/실행 중 {len(game['pending'])}, "
              f"신규 {game['new_seeds']} -> 완료 예정 {game['expected_seed_count']}개")
    for job in report["new_jobs"]:
        gap = f", |ΔHNS|={job['absolute_hns_gap']:.6f}" if job["absolute_hns_gap"] is not None else ""
        print(f"{job['game']}-{job['seed']} ({job['round_target']}개째): {job['reason']}{gap}")
        print(job["command"])
    verb = "추가 완료" if args.apply else "추가 예정 (계획만 확인)"
    print(f"{verb}: {len(report['new_jobs'])}개 명령, {report['queue']}")
    if args.report:
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return report


if __name__ == "__main__":
    main()
