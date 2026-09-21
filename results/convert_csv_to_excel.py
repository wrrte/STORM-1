import ast
from copy import copy
from pathlib import Path

import pandas as pd
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Side, Font, PatternFill
from openpyxl.utils import get_column_letter

# (Random, Human, STORM 논문) 고정 참조 점수.
# Random/Human: iclr2027_conference.tex의 tab:main_performance (2026-09-17).
# STORM: 첨부된 논문 표의 STORM (ours) 열. 실행 시 원본 문서를 읽지 않습니다.
REFERENCE_SCORES = {
    'Alien': (227.8, 7127.7, 984),
    'Amidar': (5.8, 1719.5, 205),
    'Assault': (222.4, 742.0, 801),
    'Asterix': (210.0, 8503.3, 1028),
    'BankHeist': (14.2, 753.1, 641),
    'BattleZone': (2360.0, 37187.5, 13540),
    'Boxing': (0.1, 12.1, 80),
    'Breakout': (1.7, 30.5, 16),
    'ChopperCommand': (811.0, 7387.8, 1888),
    'CrazyClimber': (10780.5, 35829.4, 66776),
    'DemonAttack': (152.1, 1971.0, 165),
    'Freeway': (0.0, 29.6, 34),
    'Frostbite': (65.2, 4334.7, 1316),
    'Gopher': (257.6, 2412.5, 8240),
    'Hero': (1027.0, 30826.4, 11044),
    'Jamesbond': (29.0, 302.8, 509),
    'Kangaroo': (52.0, 3035.0, 4208),
    'Krull': (1598.0, 2665.5, 8413),
    'KungFuMaster': (258.5, 22736.3, 26182),
    'MsPacman': (307.3, 6951.6, 2673),
    'Pong': (-20.7, 14.6, 11),
    'PrivateEye': (24.9, 69571.3, 7781),
    'Qbert': (163.9, 13455.0, 4522),
    'RoadRunner': (11.5, 7845.0, 17564),
    'Seaquest': (68.4, 42054.7, 525),
    'UpNDown': (533.4, 11693.2, 7985),
}

PAIRED_MEAN_COLUMN = 'Mean (공통 시드)'
SCORE_DELTA_COLUMN = 'Δ Score (행별 비교)'
HNS_DELTA_COLUMN = 'Δ HNS (행별 비교)'


def expand_shared_runs(df):
    """공통 warmup의 진행 상태를 선택된 실험의 실제 config로 펼칩니다."""
    # training_branches.py의 RETRIEVAL_EXPERIMENTS와 동일한 부분 설정입니다.
    overrides = {
        'baseline': {'Retrieval Enable': False},
        'retrieval': {'Retrieval Enable': True},
        'target1': {'Retrieval Enable': True, 'Retrieval Target': 1},
        'value': {'Retrieval Enable': True, 'Value Signal': 'value'},
        'add': {'Retrieval Enable': True, 'Score Combination': 'add'},
    }
    rows = []
    for row in df.to_dict('records'):
        enable = str(row['Retrieval Enable']).strip()
        if enable.startswith('['):
            experiments = ast.literal_eval(enable)
            if not isinstance(experiments, list) or not experiments:
                raise ValueError(f'Invalid Retrieval.enable experiment list: {enable}')
            experiments = [str(name).strip().lower() for name in experiments]
        elif enable.lower() == 'both' or str(row['Run Name']).lower().endswith('_both'):
            experiments = ['baseline', 'retrieval']
        else:
            rows.append(row)
            continue

        if row['State'] != 'running':
            continue
        for name in dict.fromkeys(experiments):
            if name not in overrides:
                raise ValueError(f'Unknown Retrieval.enable experiment: {name}')
            rows.append({**row, **overrides[name]})
    return pd.DataFrame(rows, columns=list(dict.fromkeys([
        *df.columns, 'Value Signal', 'Score Combination',
    ])))


def config_value(row, column, default):
    """새 config 필드가 없는 과거 CSV/백업에는 학습 코드의 기본값을 적용합니다."""
    value = row.get(column)
    if pd.isna(value) or str(value).strip().lower() in {'', 'n/a', 'nan', 'none'}:
        return default
    return value


def load_excluded_seeds():
    """update_tex.py를 실행하지 않고 현재 EXCLUDED_SEEDS 설정을 읽습니다."""
    source_path = Path(__file__).resolve().with_name('update_tex.py')
    tree = ast.parse(source_path.read_text(encoding='utf-8'), filename=str(source_path))
    for node in tree.body:
        if isinstance(node, ast.Assign):
            targets = node.targets
        elif isinstance(node, ast.AnnAssign):
            targets = [node.target]
        else:
            continue
        if any(isinstance(target, ast.Name) and target.id == 'EXCLUDED_SEEDS' for target in targets):
            return ast.literal_eval(node.value)
    raise ValueError(f'EXCLUDED_SEEDS 설정을 찾을 수 없습니다: {source_path}')


def mark_excluded_seeds(worksheet, pivot_df, start_row, excluded_seeds):
    """점수와 시드 열 이름을 보존하면서 제외된 실험 셀의 서식만 변경합니다."""
    excluded_fill = PatternFill(fill_type='solid', fgColor='E7E6E6')
    for row_idx, (game, config) in enumerate(pivot_df.index):
        if config.startswith('STORM (논문)'):
            continue
        for seed in excluded_seeds.get(game, set()):
            if seed not in pivot_df.columns:
                continue
            cell = worksheet.cell(
                row=start_row + row_idx,
                column=pivot_df.columns.get_loc(seed) + 3,
            )
            font = copy(cell.font)
            font.strike = True
            if 'RUNNING' not in str(cell.value):
                cell.fill = excluded_fill
                font.color = '808080'
            cell.font = font
            cell.comment = Comment(
                f'EXCLUDED_SEEDS: {game}, seed {seed}\n'
                'STORM/results/update_tex.py의 EXCLUDED_SEEDS에 지정되어 '
                '공통 시드 평균, Δ Score, Δ HNS 및 LaTeX 결과 집계에서 제외되는 시드입니다.\n'
                '취소선은 제외된 시드, 노란색 배경은 실행 중인 run을 뜻합니다.',
                'STORM',
            )


def parse_score(value):
    """update_tex.py의 parse_val과 동일하게 셀의 첫 번째 점수만 사용."""
    value = str(value).strip()
    if value.lower() in {'', 'nan', 'n/a', 'na', 'none', 'running'}:
        return float('nan')
    value = value.split(',')[0].split('(')[0].strip()
    try:
        return float(value)
    except (TypeError, ValueError):
        return float('nan')


def mark_save_warmup(worksheet, pivot_df, start_row, highlighted_cells):
    """config 또는 실행 인자에서 save_warmup=True인 run의 시드 셀을 강조합니다."""
    warmup_fill = PatternFill(fill_type='solid', fgColor='C6EFCE')
    warmup_side = Side(border_style='medium', color='00B050')
    for game, config, seed in highlighted_cells:
        cell = worksheet.cell(
            row=start_row + pivot_df.index.get_loc((game, config)),
            column=pivot_df.columns.get_loc(seed) + 3,
        )
        font = copy(cell.font)
        font.bold = True
        # 실행 중/제외 시드의 배경색과 취소선을 보존하고 초록 테두리로 함께 표시합니다.
        if 'RUNNING' not in str(cell.value) and not font.strike:
            cell.fill = warmup_fill
            font.color = '006100'
        cell.font = font
        border = copy(cell.border)
        border.left = border.right = border.top = border.bottom = warmup_side
        cell.border = border
        note = (
            '초록색 테두리: 이 셀에 표시된 run 중 config 또는 실행 인자에서 '
            'JointTrainAgent.Retrieval.save_warmup=True인 run이 있습니다. '
            '자식 분기는 재저장을 방지하기 위해 최종 config가 False일 수 있습니다.'
        )
        if cell.comment:
            cell.comment.text += '\n\n' + note
        else:
            cell.comment = Comment(note, 'STORM')


def add_paired_baseline_mean(pivot_df, excluded_seeds):
    """제외 목록을 반영한 공통 시드 평균과 논문/target 16 행의 baseline 대비 비교를 추가."""
    seed_columns = [column for column in pivot_df.columns if str(column).strip().isdigit()]
    columns = list(pivot_df.columns)
    columns += [
        '   ', PAIRED_MEAN_COLUMN, SCORE_DELTA_COLUMN, HNS_DELTA_COLUMN,
    ]
    pivot_df = pivot_df.reindex(columns=columns, fill_value='')

    for game in pivot_df.index.get_level_values('Game').unique():
        baseline_key = (game, 'Retrieval 미사용')
        target_key = (game, 'target: 16 (anchor 미설정)')
        if baseline_key not in pivot_df.index or target_key not in pivot_df.index:
            continue
        game_seed_columns = [
            column for column in seed_columns
            if int(str(column).strip()) not in excluded_seeds.get(game, set())
        ]
        baseline = pivot_df.loc[baseline_key, game_seed_columns].map(parse_score)
        target = pivot_df.loc[target_key, game_seed_columns].map(parse_score)
        valid = baseline.notna() & target.notna()
        if valid.any():
            baseline_mean = baseline[valid].mean()
            target_mean = target[valid].mean()
            target_delta = target_mean - baseline_mean
            pivot_df.at[baseline_key, PAIRED_MEAN_COLUMN] = baseline_mean
            pivot_df.at[target_key, PAIRED_MEAN_COLUMN] = target_mean
            pivot_df.at[target_key, SCORE_DELTA_COLUMN] = abs(target_delta)
            if game in REFERENCE_SCORES:
                random_score, human_score, paper_score = REFERENCE_SCORES[game]
                paper_key = (game, f'STORM (논문): {paper_score}')
                score_delta = baseline_mean - paper_score
                if paper_key in pivot_df.index:
                    pivot_df.at[paper_key, SCORE_DELTA_COLUMN] = score_delta
                if human_score != random_score:
                    # HNS는 Random=0, Human=1 기준의 정규화 점수이며 백분율이 아닙니다.
                    pivot_df.at[target_key, HNS_DELTA_COLUMN] = (
                        target_delta / (human_score - random_score)
                    )
                    if paper_key in pivot_df.index:
                        pivot_df.at[paper_key, HNS_DELTA_COLUMN] = (
                            score_delta / (human_score - random_score)
                        )
    return pivot_df


def main():
    excluded_seeds = load_excluded_seeds()
    # Load CSV
    df = pd.read_csv('wandb_runs_classification.csv')
    if 'State' not in df:
        df['State'] = ''
    df['State'] = df['State'].fillna('').astype(str).str.strip().str.lower()

    df = expand_shared_runs(df)

    # Data extraction
    data = []
    for idx, row in df.iterrows():
        run_name = str(row['Run Name'])
        parts = run_name.split('_')
        game = parts[0]
        
        # Extract seed
        seed = str(row['Seed'])
        if seed.endswith('.0'):
            seed = seed[:-2]
        if seed == 'N/A' or seed == 'nan' or seed == 'None':
            # Try to extract from Run Name (e.g., Amidar_f05wv0rj_2000_X)
            if len(parts) >= 3 and parts[2].isdigit():
                seed = parts[2]
            elif len(parts) == 4 and parts[3].isdigit():
                seed = parts[3]
            
        if not str(seed).isdigit():
            continue # Ignore if seed cannot be found
            
        seed = int(seed)
        
        ret_enable = str(row['Retrieval Enable']).strip().lower() in ['true', '1', 't']
        created_at_utc = pd.to_datetime(row.get('Created At', '2000-01-01T00:00:00Z'), utc=True)
        
        # ---------------------------------------------------------
        # 날짜 필터링: EMA 기법이 최신 코드에 적용된 시점 이후의 실험만 포함
        # 기준 시점: 2026-08-24 11:05:43 UTC
        #
        # [AGENT INSTRUCTION] 
        # Retrieval 미사용(ret_enable == False) 설정은 이 날짜 필터링을 무시하고
        # 모든 시점의 데이터를 가져오며, 중복 시 최신 점수를 기록하도록 항상 이 로직을 유지해야 합니다.
        # 향후 이 코드를 수정할 계획을 세울 때, AI 에이전트는 반드시 사용자에게 
        # "Retrieval 미사용은 예전 데이터를 유지하는 로직이 적용되어 있다"는 사실을 사전에 알려야 합니다.
        # ---------------------------------------------------------
        ema_applied_date = pd.to_datetime("2026-08-24T11:05:43Z", utc=True)
        if ret_enable and created_at_utc < ema_applied_date:
            continue
        
        # Safely parse warmup_steps to int if possible
        try:
            w_float = float(row['Warmup Steps'])
            if not pd.isna(w_float):
                warmup_steps = int(w_float)
            else:
                warmup_steps = 'N/A'
        except (ValueError, TypeError):
            warmup_steps = 'N/A'

        try:
            cw_float = float(row.get('Calculated Warmup Steps', row['Warmup Steps']))
            if not pd.isna(cw_float):
                calculated_warmup_steps = int(cw_float)
            else:
                calculated_warmup_steps = 'N/A'
        except (ValueError, TypeError, KeyError):
            calculated_warmup_steps = 'N/A'

        # ---------------------------------------------------------
        # 실험 필터링 및 표기법 변경 로직
        # 1. Retrieval을 사용하지 않은 실험은 'Retrieval 미사용'으로 표기
        # 2. Warmup 50000, BSR retrieved, 적용 임계값(multiply: 3.5 / add: 4.0)만 유지
        # 3. target 1 또는 16이고 anchor 미설정인 실험만 유지 (기존대로 0.0625도 미설정 취급)
        # 4. value_signal=value / score_combination=add는 별도 config 행으로 표시
        # ---------------------------------------------------------
        if not ret_enable:
            config = 'Retrieval 미사용'
        else:
            bsr = row.get('Batch Size Reduction', 'N/A')
            if pd.isna(bsr) or str(bsr).strip() == '' or str(bsr) == 'nan':
                bsr = 'N/A'

            value_signal = str(config_value(row, 'Value Signal', 'value_diff')).strip().lower()
            score_combination = str(config_value(row, 'Score Combination', 'multiply')).strip().lower()
            if value_signal not in ('value_diff', 'value') or score_combination not in ('multiply', 'add'):
                continue
            if score_combination == 'add':
                z_score = config_value(row, 'Additive Z Score Threshold', 4.0)
                expected_z_score = 4.0
            else:
                z_score = row.get('Z Score Threshold', 'N/A')
                expected_z_score = 3.5
            try:
                z_score_float = float(z_score)
            except (ValueError, TypeError):
                z_score_float = None

            if warmup_steps == 50000 and str(bsr) == 'retrieved' and z_score_float == expected_z_score:
                r_target = row.get('Retrieval Target', 'N/A')
                try:
                    r_target = float(r_target)
                    if r_target not in (1, 16):
                        continue
                except (ValueError, TypeError):
                    continue
                
                a_weight = row.get('Anchor Weight', 'N/A')
                try:
                    anchor_is_unset = float(a_weight) == 0.0625
                except (ValueError, TypeError):
                    anchor_is_unset = False

                if (pd.isna(a_weight) or str(a_weight).strip() == '' or
                        str(a_weight) == 'nan' or str(a_weight) == 'N/A' or
                        anchor_is_unset):
                    config = f'target: {int(r_target)} (anchor 미설정)'
                    variants = []
                    if value_signal == 'value':
                        variants.append('value')
                    if score_combination == 'add':
                        variants.append('add')
                    if variants:
                        config += f" [{', '.join(variants)}]"
                else:
                    continue
            else:
                # 위 경우에 해당하지 않는 실험은 제외
                continue
                
        eval_return = row['Eval Return']
        h_bits = row.get('Hash Bits', 'N/A')
        if pd.isna(h_bits) or str(h_bits).strip() == '' or str(h_bits) == 'nan':
            h_bits = 'N/A'

        data.append({
            'Game': game,
            'Config': config,
            'Seed': seed,
            'State': row['State'],
            'Eval Return': eval_return,
            'Warmup Steps': calculated_warmup_steps,
            'Hash Bits': h_bits,
            'Save Warmup': any(
                str(row.get(column, '')).strip().lower() == 'true'
                for column in ('Save Warmup', 'Save Warmup Requested')
            ),
            'Created At': created_at_utc
        })

    # 강제로 Frostbite - Retrieval 미사용의 누락된 시드 추가
    data.append({
        'Game': 'Frostbite',
        'Config': 'Retrieval 미사용',
        'Seed': 10,
        'State': 'finished',
        'Eval Return': '2068',
        'Warmup Steps': 'N/A',
        'Hash Bits': 'N/A',
        'Save Warmup': False,
        'Created At': pd.to_datetime("2026-08-11T04:36:29Z", utc=True)
    })
    data.append({
        'Game': 'Frostbite',
        'Config': 'Retrieval 미사용',
        'Seed': 3710,
        'State': 'finished',
        'Eval Return': '1904',
        'Warmup Steps': 'N/A',
        'Hash Bits': 'N/A',
        'Save Warmup': False,
        'Created At': pd.to_datetime("2026-08-11T04:36:35Z", utc=True)
    })

    parsed_df = pd.DataFrame(data)

    # 1. 실행 중인 run은 평가 점수가 없어도 유지합니다.
    parsed_df['Eval Return'] = parsed_df['Eval Return'].astype(str)
    parsed_df = parsed_df[
        parsed_df['State'].eq('running') | (
            (parsed_df['Eval Return'] != 'N/A') &
            (parsed_df['Eval Return'] != 'nan') &
            (parsed_df['Eval Return'].str.strip() != '')
        )
    ]
    
    # 2. 모든 실험 결과 표시 (중복 런 포함 모두 나열하기 위해 중복 제거 로직 삭제)
    # 먼저 Created At 기준으로 내림차순 정렬 (최신이 위로 오도록)
    parsed_df = parsed_df.sort_values(by='Created At', ascending=False)
    
    # 3. 동일한 시드에서 여러 결과가 있을 경우, 셀 하나에 여러 줄로 나열
    def aggregate_cell(group):
        # 점수와 강조 여부 모두 실제로 표시되는 run만 기준으로 계산합니다.
        is_hb_10 = group['Hash Bits'].apply(lambda x: str(x).strip() in ['10', '10.0'])
        if is_hb_10.any():
            group = group[is_hb_10 | group['State'].eq('running')]

        def format_single_row(r):
            val = 'RUNNING' if r['State'] == 'running' else r['Eval Return']
            try:
                val = f"{float(val):.2f}"
            except ValueError:
                pass
                
            # 단일 결과여도 50000이 아니면 명시
            warmup = r['Warmup Steps']
            try:
                w_val = int(float(warmup))
                if w_val != 50000:
                    val = f"{val} (w: {w_val})"
            except (ValueError, TypeError):
                pass
                
            return str(val)

        if len(group) == 1:
            value = format_single_row(group.iloc[0])
        else:
            items = []
            has_diff_hash = group['Hash Bits'].nunique() > 1
            has_diff_warmup = group['Warmup Steps'].nunique() > 1
            for _, r in group.iterrows():
                val = 'RUNNING' if r['State'] == 'running' else r['Eval Return']
                try:
                    val = f"{float(val):.2f}"
                except ValueError:
                    pass
                
                extras = []
                if has_diff_hash:
                    extras.append(f"hb: {r['Hash Bits']}")
                
                warmup = r['Warmup Steps']
                # 값이 다를 경우 구분을 위해 무조건 덧붙여 표기
                if has_diff_warmup or (str(warmup) != '50000' and str(warmup) != 'N/A'):
                    extras.append(f"w: {warmup}")
                
                if extras:
                    items.append(f"{val} ({', '.join(extras)})")
                else:
                    items.append(f"{val}")
            value = ", ".join(items)

        return pd.Series({
            'Final Eval Return': value,
            'Save Warmup': group['Save Warmup'].any(),
        })

    agg_df = parsed_df.groupby(['Game', 'Config', 'Seed']).apply(aggregate_cell, include_groups=False).reset_index()
    highlighted_cells = set(
        agg_df.loc[agg_df['Save Warmup'], ['Game', 'Config', 'Seed']]
        .itertuples(index=False, name=None)
    )
    
    # Pivot table
    pivot_df = agg_df.pivot_table(
        index=['Game', 'Config'],
        columns='Seed',
        values='Final Eval Return',
        aggfunc='first'
    )

    # 시드 그룹 순서. 6020은 결과가 없어도 60*0 그룹 내에 표시합니다.
    target_columns = [1, 2, 10, 710, 1710, 2710, 3710, ' ', 2000, 2010, '  ', 5090]
    extra_seeds = sorted((set(pivot_df.columns) | {6020}) - set(target_columns))
    seeds_999 = [seed for seed in extra_seeds if str(seed).startswith('999')]
    full_columns = target_columns + [seed for seed in extra_seeds if seed not in seeds_999]
    if seeds_999:
        full_columns += ['    '] + seeds_999
    
    pivot_df = pivot_df.reindex(columns=full_columns)
    pivot_df = pivot_df.fillna('')

    # 논문 점수는 시드별 실험값과 구분하여 Config 열의 별도 첫 행에 표시합니다.
    reference_index = pd.MultiIndex.from_tuples(
        [(game, f'STORM (논문): {REFERENCE_SCORES[game][2]}')
         for game in pivot_df.index.get_level_values('Game').unique()
         if game in REFERENCE_SCORES],
        names=['Game', 'Config'],
    )
    reference_df = pd.DataFrame('', index=reference_index, columns=pivot_df.columns)
    pivot_df = pd.concat([pivot_df, reference_df])

    # Sorting Index
    base_order = [
        'STORM (논문)',
        'Retrieval 미사용',
        'target'
    ]

    def get_sort_key(config_str):
        variant_order = next((
            order for order, suffix in enumerate([' [value]', ' [add]', ' [value, add]'], start=1)
            if config_str.endswith(suffix)
        ), 0)
        for i, base in enumerate(base_order):
            if config_str.startswith(base):
                return i, variant_order, config_str
        return len(base_order), variant_order, config_str

    # Sort the multi-index: alphabetical by Game, then by specified Config order
    sorted_index = sorted(pivot_df.index, key=lambda x: (x[0], get_sort_key(x[1])))
    pivot_df = pivot_df.reindex(sorted_index)
    pivot_df = add_paired_baseline_mean(pivot_df, excluded_seeds)

    output_path = 'converted_results.xlsx'
    
    with pd.ExcelWriter(output_path, engine='openpyxl') as writer:
        pivot_df.to_excel(writer, sheet_name='Results')
        
        workbook = writer.book
        worksheet = writer.sheets['Results']
        # 세로 스크롤 시 첫 행의 시드 번호 헤더를 고정합니다.
        worksheet.freeze_panes = 'A2'

        mean_col = pivot_df.columns.get_loc(PAIRED_MEAN_COLUMN) + 3
        for column_idx, column in enumerate(pivot_df.columns, start=3):
            if not str(column).strip():
                worksheet.column_dimensions[get_column_letter(column_idx)].width = 3
        for offset, (width, number_format) in enumerate([
            (30, '0.00'),
            (46, '+0.00;-0.00;0.00'),
            (32, '+0.0000;-0.0000;0.0000'),
        ]):
            column = mean_col + offset
            worksheet.column_dimensions[get_column_letter(column)].width = width
            for row in worksheet.iter_rows(min_row=2, min_col=column, max_col=column):
                row[0].number_format = number_format

        worksheet.cell(row=1, column=mean_col).comment = Comment(
            'Retrieval 미사용과 target: 16 양쪽에 유효한 점수가 있는 공통 시드만 사용합니다. '
            'update_tex.py의 EXCLUDED_SEEDS에 지정된 시드는 제외합니다. '
            '각 시드에 여러 결과가 있으면 첫 번째 점수를 사용합니다.',
            'STORM',
        )
        worksheet.cell(row=1, column=mean_col + 1).comment = Comment(
            'STORM (논문) 행: Retrieval 미사용의 공통 시드 평균 − STORM 논문 점수.\n'
            'target: 16 행: |target: 16의 공통 시드 평균 − Retrieval 미사용의 공통 시드 평균|.',
            'STORM',
        )
        worksheet.cell(row=1, column=mean_col + 2).comment = Comment(
            'STORM (논문) 행: (Retrieval 미사용 평균 − STORM 논문 점수) / (Human − Random).\n'
            'target: 16 행: (target: 16 평균 − Retrieval 미사용 평균) / (Human − Random).\n'
            '모든 평균은 EXCLUDED_SEEDS를 제외한 공통 시드 기준이며 '
            'HNS 차이는 부호를 유지합니다. 백분율이 아닙니다.',
            'STORM',
        )

        start_row = worksheet.max_row - len(pivot_df) + 1
        game_row_counts = pivot_df.groupby(level='Game').size()
        for row_idx, (game, config) in enumerate(pivot_df.index):
            excel_row = row_idx + start_row
            if config == 'target: 16 (anchor 미설정)':
                worksheet.cell(row=excel_row, column=mean_col + 1).number_format = '0.00'
            # 병합된 게임 셀의 세 줄이 모두 보이도록 높이와 줄바꿈 설정.
            worksheet.row_dimensions[excel_row].height = max(24, 72 / game_row_counts[game])
            if row_idx == 0 or game != pivot_df.index[row_idx - 1][0]:
                if game in REFERENCE_SCORES:
                    random_score, human_score, _ = REFERENCE_SCORES[game]
                    cell = worksheet.cell(row=excel_row, column=1)
                    cell.value = f'{game}\nRandom: {random_score:.1f}\nHuman: {human_score:.1f}'
                    cell.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
        
        # 1. 모든 셀 폰트 크기 16으로 설정
        for row in worksheet.iter_rows():
            for cell in row:
                if cell.font:
                    cell.font = Font(name=cell.font.name, size=16, bold=cell.font.bold, italic=cell.font.italic, color=cell.font.color)
                else:
                    cell.font = Font(size=16)

        # Drama와 동일한 색상으로 실행 중인 run이 포함된 시드 셀을 강조합니다.
        running_fill = PatternFill(fill_type='solid', fgColor='FFF2CC')
        for column_idx, column in enumerate(pivot_df.columns, start=3):
            if not str(column).strip().isdigit():
                continue
            for row in worksheet.iter_rows(min_row=start_row, min_col=column_idx, max_col=column_idx):
                cell = row[0]
                if 'RUNNING' in str(cell.value):
                    cell.fill = running_fill
                    cell.font = Font(name=cell.font.name, size=16, bold=True, color='9C6500')

        mark_excluded_seeds(worksheet, pivot_df, start_row, excluded_seeds)
                    
        # 2. A, B열 내용에 맞게 열 너비 자동 맞춤
        for col_letter, col_idx in [('A', 1), ('B', 2)]:
            max_len = 0
            for row in worksheet.iter_rows(min_col=col_idx, max_col=col_idx):
                for cell in row:
                    if cell.value:
                        # 줄바꿈이 있는 경우 가장 긴 줄 기준
                        lines = str(cell.value).split('\n')
                        for line in lines:
                            max_len = max(max_len, len(line))
                            
            # 폰트 사이즈 16이므로, 일반(11)보다 크기 때문에 가중치(약 1.7배)를 줍니다.
            worksheet.column_dimensions[col_letter].width = (max_len * 1.7) + 2

        # medium border for visibility between games
        top_side = Side(border_style="medium", color="000000")
        
        # Calculate start_row dynamically
        start_row = worksheet.max_row - len(pivot_df) + 1
        
        for row_idx, (game, config) in enumerate(pivot_df.index):
            excel_row = row_idx + start_row
            
            # Draw line above the first row of each new game
            if row_idx == 0 or pivot_df.index[row_idx][0] != pivot_df.index[row_idx-1][0]:
                for col_idx in range(1, len(pivot_df.columns) + 3):
                    cell = worksheet.cell(row=excel_row, column=col_idx)
                    cell.border = Border(
                        top=top_side,
                        left=cell.border.left,
                        right=cell.border.right,
                        bottom=cell.border.bottom
                    )

        mark_save_warmup(worksheet, pivot_df, start_row, highlighted_cells)

    print(f"Successfully saved to {output_path}")

if __name__ == '__main__':
    main()
